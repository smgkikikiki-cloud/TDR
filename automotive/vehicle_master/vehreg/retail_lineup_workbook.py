"""Excel contract for the owner-authoritative Retail Lineup Bootstrap.

Chunk 3 is deliberately compile-only. A workbook can describe the complete
target CURRENT lineup and be compiled into the immutable Chunk 1 plan, but this
module never calls the Chunk 2 apply engine and never writes canonical data.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timezone
import hashlib
import json
from pathlib import Path
import re
from typing import Any, Iterable, Mapping, Sequence

from .catalog import Catalog, DATA_DIR, DEFAULT_YEAR
from .current_retail import load_current_retail_index
from .retail_lifecycle_review import load_trim_lifecycle_decisions
from .retail_lineup_bootstrap import (
    RetailLineupBootstrapError,
    RetailLineupPlan,
    _decision_index,
    _resolved_current,
    baseline_hash_for_models,
    plan_retail_lineup,
)

SCHEMA_VERSION = 1
MODE = "REPLACE_AND_ARCHIVE"
TARGET_SHEET = "TARGET_LINEUP"
SNAPSHOT_SHEET = "CURRENT_SNAPSHOT"
META_SHEET = "IMPORT_META"

TARGET_COLUMNS = (
    "model_id",
    "generation_id",
    "trim_name",
    "powertrain",
    "canonical_trim_id",
    "notes",
)
REQUIRED_TARGET_COLUMNS = TARGET_COLUMNS[:4]
SNAPSHOT_COLUMNS = (
    "brand",
    "model",
    "model_id",
    "generation_id",
    "canonical_trim_id",
    "trim_name",
    "powertrain",
    "model_retail_status",
    "generation_ended",
    "lifecycle_decision",
    "approved_current",
    "resolved_current",
    "price_records",
    "spec_facts",
    "campaign_price_records",
)
META_KEYS = (
    "schema_version",
    "generated_at",
    "as_of",
    "catalog_year",
    "base_release_id",
    "baseline_hash",
    "mode",
    "target_model_ids",
)
_HEX64 = re.compile(r"^[0-9a-fA-F]{64}$")


class RetailLineupWorkbookError(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class RetailLineupWorkbookMeta:
    schema_version: int
    generated_at: str
    as_of: str
    catalog_year: int
    base_release_id: str
    baseline_hash: str
    mode: str
    target_model_ids: tuple[str, ...]

    def as_dict(self) -> dict[str, Any]:
        return {
            "schema_version": self.schema_version,
            "generated_at": self.generated_at,
            "as_of": self.as_of,
            "catalog_year": self.catalog_year,
            "base_release_id": self.base_release_id,
            "baseline_hash": self.baseline_hash,
            "mode": self.mode,
            "target_model_ids": list(self.target_model_ids),
        }


@dataclass(frozen=True, slots=True)
class RetailLineupWorkbookCompileResult:
    source_sha256: str
    meta: RetailLineupWorkbookMeta
    rows_read: int
    plan: RetailLineupPlan

    def as_dict(self) -> dict[str, Any]:
        return {
            "schema_version": 1,
            "source_sha256": self.source_sha256,
            "meta": self.meta.as_dict(),
            "rows_read": self.rows_read,
            "compile_only": True,
            "plan": self.plan.as_dict(),
        }


def _value(raw: Any) -> str:
    if raw is None:
        return ""
    if isinstance(raw, bool):
        return "TRUE" if raw else "FALSE"
    return str(raw).strip()


def _enum_value(raw: Any) -> str:
    return str(getattr(raw, "value", raw or "")).strip()


def _generated_at(raw: str | None) -> str:
    value = str(raw or "").strip()
    if not value:
        return datetime.now(timezone.utc).isoformat(timespec="seconds")
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError as exc:
        raise RetailLineupWorkbookError("generated_at must be ISO-8601") from exc
    if parsed.tzinfo is None:
        raise RetailLineupWorkbookError("generated_at must include timezone")
    return parsed.isoformat(timespec="seconds")


def _as_of(raw: str | None, generated_at: str) -> str:
    value = str(raw or "").strip() or datetime.fromisoformat(generated_at).date().isoformat()
    try:
        return date.fromisoformat(value).isoformat()
    except ValueError as exc:
        raise RetailLineupWorkbookError("as_of must be YYYY-MM-DD") from exc


def _model_ids(values: Iterable[str]) -> tuple[str, ...]:
    ids = tuple(sorted({_value(value) for value in values if _value(value)}))
    if not ids:
        raise RetailLineupWorkbookError("at least one target model_id is required")
    return ids


def _read_json(path: Path) -> dict[str, Any]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise RetailLineupWorkbookError(f"cannot read snapshot source {path}: {exc}") from exc
    if not isinstance(payload, dict):
        raise RetailLineupWorkbookError(f"snapshot source {path} must be a JSON object")
    return payload


def _presence_indexes(data_dir: Path, year: int) -> tuple[dict[str, int], dict[str, int], dict[str, int]]:
    """Informational-only counts for CURRENT_SNAPSHOT.

    The snapshot is never read back into the write plan, so these are intentionally
    simple presence counts over the canonical stores rather than a second resolver.
    """
    price_counts: dict[str, int] = {}
    campaign_counts: dict[str, int] = {}
    spec_counts: dict[str, int] = {}

    price_root = data_dir / str(year) / "market" / "prices"
    if price_root.is_dir():
        for path in sorted(price_root.glob("*.json")):
            payload = _read_json(path)
            rows = payload.get("prices") or []
            if not isinstance(rows, list):
                raise RetailLineupWorkbookError(f"{path}: prices must be an array")
            for row in rows:
                if not isinstance(row, dict):
                    continue
                trim_id = _value(row.get("trim_id"))
                if not trim_id:
                    continue
                price_counts[trim_id] = price_counts.get(trim_id, 0) + 1
                if _value(row.get("campaign_id")):
                    campaign_counts[trim_id] = campaign_counts.get(trim_id, 0) + 1

    spec_root = data_dir / str(year) / "product" / "comparable_specs" / "facts"
    if spec_root.is_dir():
        for path in sorted(spec_root.glob("*.json")):
            payload = _read_json(path)
            rows = payload.get("facts") or []
            if not isinstance(rows, list):
                raise RetailLineupWorkbookError(f"{path}: facts must be an array")
            for row in rows:
                if not isinstance(row, dict):
                    continue
                trim_id = _value(row.get("trim_id"))
                if trim_id:
                    spec_counts[trim_id] = spec_counts.get(trim_id, 0) + 1

    return price_counts, spec_counts, campaign_counts


def _generation_model_id(catalog: Catalog, generation_id: str) -> str:
    generation = catalog.generations.get(generation_id)
    return _value(getattr(generation, "model_id", "")) if generation else ""


def _trim_model_id(catalog: Catalog, trim_id: str) -> str:
    trim = catalog.trims.get(trim_id)
    return _generation_model_id(catalog, _value(getattr(trim, "generation_id", ""))) if trim else ""


def _resolved_for_model(
    catalog: Catalog,
    model_id: str,
    current: Mapping[str, Iterable[str]],
    decisions: Mapping[str, str],
    as_of: str,
) -> tuple[str, ...]:
    return _resolved_current(catalog, model_id, current, decisions, date.fromisoformat(as_of))


def _active_generation_ids(catalog: Catalog, model_id: str, as_of: str) -> tuple[str, ...]:
    day = date.fromisoformat(as_of)
    out = []
    for generation_id, generation in catalog.generations.items():
        if _value(getattr(generation, "model_id", "")) != model_id:
            continue
        ended = _value(getattr(generation, "ended", ""))
        if ended and date.fromisoformat(ended) <= day:
            continue
        out.append(generation_id)
    return tuple(sorted(out))


def _header(ws) -> tuple[str, ...]:
    values = [_value(ws.cell(1, column).value) for column in range(1, ws.max_column + 1)]
    while values and not values[-1]:
        values.pop()
    return tuple(values)


def _reject_formula(cell, *, location: str) -> None:
    if getattr(cell, "data_type", "") == "f":
        raise RetailLineupWorkbookError(f"{location}: formulas are not allowed")


def _read_meta(ws) -> RetailLineupWorkbookMeta:
    if _header(ws) != ("key", "value"):
        raise RetailLineupWorkbookError("IMPORT_META headers must be exactly key,value")
    raw: dict[str, Any] = {}
    for row in range(2, ws.max_row + 1):
        key_cell, value_cell = ws.cell(row, 1), ws.cell(row, 2)
        _reject_formula(key_cell, location=f"IMPORT_META!A{row}")
        _reject_formula(value_cell, location=f"IMPORT_META!B{row}")
        key = _value(key_cell.value)
        value = value_cell.value
        if not key and value in (None, ""):
            continue
        if key not in META_KEYS:
            raise RetailLineupWorkbookError(f"IMPORT_META unknown key {key!r}")
        if key in raw:
            raise RetailLineupWorkbookError(f"IMPORT_META duplicate key {key!r}")
        raw[key] = value
    missing = [key for key in META_KEYS if key not in raw]
    if missing:
        raise RetailLineupWorkbookError(f"IMPORT_META missing keys: {missing}")

    try:
        schema_version = int(raw["schema_version"])
        catalog_year = int(raw["catalog_year"])
    except (TypeError, ValueError) as exc:
        raise RetailLineupWorkbookError("schema_version and catalog_year must be integers") from exc
    if schema_version != SCHEMA_VERSION:
        raise RetailLineupWorkbookError(f"unsupported schema_version {schema_version}")

    generated = _generated_at(_value(raw["generated_at"]))
    as_of = _as_of(_value(raw["as_of"]), generated)
    base_release_id = _value(raw["base_release_id"])
    if not base_release_id:
        raise RetailLineupWorkbookError("base_release_id is required")
    baseline_hash = _value(raw["baseline_hash"]).lower()
    if not _HEX64.fullmatch(baseline_hash):
        raise RetailLineupWorkbookError("baseline_hash must be a SHA-256 hex digest")
    mode = _value(raw["mode"]).upper()
    if mode != MODE:
        raise RetailLineupWorkbookError(f"mode must be {MODE}")

    try:
        parsed_models = json.loads(_value(raw["target_model_ids"]))
    except json.JSONDecodeError as exc:
        raise RetailLineupWorkbookError("target_model_ids must be a JSON array") from exc
    if not isinstance(parsed_models, list) or not all(
            isinstance(value, str) and value.strip() for value in parsed_models):
        raise RetailLineupWorkbookError("target_model_ids must be a JSON array of nonempty strings")
    model_ids = _model_ids(parsed_models)
    if len(model_ids) != len(parsed_models):
        raise RetailLineupWorkbookError("target_model_ids must be unique")

    return RetailLineupWorkbookMeta(
        schema_version, generated, as_of, catalog_year, base_release_id,
        baseline_hash, mode, model_ids,
    )


def _read_target_rows(ws, target_model_ids: Sequence[str]) -> list[dict[str, str]]:
    if _header(ws) != TARGET_COLUMNS:
        raise RetailLineupWorkbookError(
            f"TARGET_LINEUP headers must be exactly {list(TARGET_COLUMNS)}")
    allowed_models = set(target_model_ids)
    rows: list[dict[str, str]] = []
    seen_models: set[str] = set()
    for row_number in range(2, ws.max_row + 1):
        cells = [ws.cell(row_number, column) for column in range(1, len(TARGET_COLUMNS) + 1)]
        for column, cell in enumerate(cells, start=1):
            _reject_formula(cell, location=f"TARGET_LINEUP!R{row_number}C{column}")
        values = [_value(cell.value) for cell in cells]
        if not any(values):
            continue
        record = dict(zip(TARGET_COLUMNS, values))
        missing = [key for key in REQUIRED_TARGET_COLUMNS if not record[key]]
        if missing:
            raise RetailLineupWorkbookError(
                f"TARGET_LINEUP row {row_number} missing required values: {missing}")
        if record["model_id"] not in allowed_models:
            raise RetailLineupWorkbookError(
                f"TARGET_LINEUP row {row_number} model_id {record['model_id']!r} "
                "is not declared in IMPORT_META target_model_ids")
        seen_models.add(record["model_id"])
        rows.append(record)

    missing_models = sorted(allowed_models - seen_models)
    if missing_models:
        raise RetailLineupWorkbookError(
            "every targeted model requires at least one target trim; "
            f"missing rows for {missing_models}. Empty target never means withdraw model.")
    return rows


def read_retail_lineup_workbook(path: Path | str) -> tuple[RetailLineupWorkbookMeta, list[dict[str, str]]]:
    """Read the three-sheet contract. CURRENT_SNAPSHOT is intentionally ignored."""
    from openpyxl import load_workbook

    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise RetailLineupWorkbookError("Retail Lineup Bootstrap accepts .xlsx only")
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise RetailLineupWorkbookError(f"cannot read workbook {path}: {exc}") from exc

    required = {TARGET_SHEET, SNAPSHOT_SHEET, META_SHEET}
    missing = sorted(required - set(workbook.sheetnames))
    if missing:
        raise RetailLineupWorkbookError(f"workbook missing required sheets: {missing}")

    meta = _read_meta(workbook[META_SHEET])
    rows = _read_target_rows(workbook[TARGET_SHEET], meta.target_model_ids)
    return meta, rows


def compile_retail_lineup_workbook(
    path: Path | str,
    *,
    data_dir: Path | str = DATA_DIR,
    expected_year: int | None = None,
) -> RetailLineupWorkbookCompileResult:
    """Compile workbook -> immutable plan only. Never calls the apply engine."""
    path = Path(path)
    meta, rows = read_retail_lineup_workbook(path)
    if expected_year is not None and int(expected_year) != meta.catalog_year:
        raise RetailLineupWorkbookError(
            f"workbook catalog_year {meta.catalog_year} != expected {expected_year}")

    root = Path(data_dir)
    catalog = Catalog.load(root, meta.catalog_year)
    current = load_current_retail_index(data_dir=root, year=meta.catalog_year)
    decisions = load_trim_lifecycle_decisions(data_dir=root, year=meta.catalog_year)

    unknown_models = sorted(set(meta.target_model_ids) - set(catalog.models))
    if unknown_models:
        raise RetailLineupWorkbookError(f"unknown target model_ids: {unknown_models}")

    actual_baseline = baseline_hash_for_models(
        catalog,
        meta.target_model_ids,
        current_retail_index=current,
        lifecycle_decisions=decisions,
        as_of=meta.as_of,
    )
    if actual_baseline != meta.baseline_hash:
        raise RetailLineupBootstrapError(
            f"STALE_BASELINE: workbook {meta.baseline_hash} != current {actual_baseline}")

    plan = plan_retail_lineup(
        rows,
        catalog=catalog,
        current_retail_index=current,
        lifecycle_decisions=decisions,
        as_of=meta.as_of,
        expected_baseline_hash=meta.baseline_hash,
        base_release_id=meta.base_release_id,
    )
    return RetailLineupWorkbookCompileResult(
        source_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        meta=meta,
        rows_read=len(rows),
        plan=plan,
    )


def generate_retail_lineup_workbook(
    destination: Path | str,
    *,
    model_ids: Sequence[str],
    base_release_id: str,
    data_dir: Path | str = DATA_DIR,
    year: int = DEFAULT_YEAR,
    generated_at: str | None = None,
    as_of: str | None = None,
) -> dict[str, Any]:
    """Generate an owner-editable workbook seeded with the resolved CURRENT lineup."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    destination = Path(destination)
    if destination.suffix.lower() != ".xlsx":
        raise RetailLineupWorkbookError("destination must end in .xlsx")
    release_id = _value(base_release_id)
    if not release_id:
        raise RetailLineupWorkbookError("base_release_id is required")
    selected = _model_ids(model_ids)
    generated = _generated_at(generated_at)
    day = _as_of(as_of, generated)
    root = Path(data_dir)

    catalog = Catalog.load(root, year)
    unknown = sorted(set(selected) - set(catalog.models))
    if unknown:
        raise RetailLineupWorkbookError(f"unknown model_ids: {unknown}")

    for model_id in selected:
        model = catalog.models[model_id]
        if _enum_value(getattr(model, "retail_status", "")).upper() == "HISTORICAL":
            raise RetailLineupWorkbookError(
                f"model {model_id} is HISTORICAL; model reactivation is outside bootstrap MVP")

    current = load_current_retail_index(data_dir=root, year=year)
    raw_decisions = load_trim_lifecycle_decisions(data_dir=root, year=year)
    decisions = _decision_index(raw_decisions)
    baseline = baseline_hash_for_models(
        catalog, selected, current_retail_index=current,
        lifecycle_decisions=raw_decisions, as_of=day)
    price_counts, spec_counts, campaign_counts = _presence_indexes(root, year)

    resolved_by_model = {
        model_id: _resolved_for_model(catalog, model_id, current, decisions, day)
        for model_id in selected
    }

    target_rows: list[list[Any]] = []
    for model_id in selected:
        trim_ids = resolved_by_model[model_id]
        if trim_ids:
            for trim_id in trim_ids:
                trim = catalog.trims[trim_id]
                target_rows.append([
                    model_id,
                    _value(trim.generation_id),
                    _value(trim.name),
                    _enum_value(trim.powertrain),
                    trim_id,
                    "",
                ])
        else:
            generations = _active_generation_ids(catalog, model_id, day)
            target_rows.append([
                model_id,
                generations[0] if len(generations) == 1 else "",
                "",
                "",
                "",
                "FILL REQUIRED: model has no resolved CURRENT trim",
            ])

    snapshot_rows: list[list[Any]] = []
    for model_id in selected:
        model = catalog.models[model_id]
        brand = catalog.brands.get(_value(model.brand_id))
        resolved = set(resolved_by_model[model_id])
        approved = set(current[model_id]) if model_id in current else None
        known_trim_ids = sorted(
            trim_id for trim_id in catalog.trims
            if _trim_model_id(catalog, trim_id) == model_id
        )
        for trim_id in known_trim_ids:
            trim = catalog.trims[trim_id]
            generation = catalog.generations.get(trim.generation_id)
            snapshot_rows.append([
                _value(getattr(brand, "name_en", "")),
                _value(getattr(model, "name_en", "")),
                model_id,
                _value(trim.generation_id),
                trim_id,
                _value(trim.name),
                _enum_value(trim.powertrain),
                _enum_value(getattr(model, "retail_status", "")),
                _value(getattr(generation, "ended", "")),
                decisions.get(trim_id, ""),
                "" if approved is None else ("YES" if trim_id in approved else "NO"),
                "YES" if trim_id in resolved else "NO",
                price_counts.get(trim_id, 0),
                spec_counts.get(trim_id, 0),
                campaign_counts.get(trim_id, 0),
            ])

    workbook = Workbook()
    target = workbook.active
    target.title = TARGET_SHEET
    snapshot = workbook.create_sheet(SNAPSHOT_SHEET)
    meta_sheet = workbook.create_sheet(META_SHEET)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for ws, headers in ((target, TARGET_COLUMNS), (snapshot, SNAPSHOT_COLUMNS), (meta_sheet, ("key", "value"))):
        ws.append(list(headers))
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        ws.freeze_panes = "A2"

    for row in target_rows:
        target.append(row)
    for row in snapshot_rows:
        snapshot.append(row)

    meta_values = {
        "schema_version": SCHEMA_VERSION,
        "generated_at": generated,
        "as_of": day,
        "catalog_year": int(year),
        "base_release_id": release_id,
        "baseline_hash": baseline,
        "mode": MODE,
        "target_model_ids": json.dumps(list(selected), ensure_ascii=False, separators=(",", ":")),
    }
    for key in META_KEYS:
        meta_sheet.append([key, meta_values[key]])

    validation = DataValidation(
        type="list",
        formula1='"ICE,HEV,PHEV,REEV,BEV,FCEV"',
        allow_blank=True,
    )
    target.add_data_validation(validation)
    validation.add("D2:D10000")

    widths = {
        TARGET_SHEET: (28, 34, 32, 14, 58, 52),
        SNAPSHOT_SHEET: (18, 24, 28, 34, 58, 30, 14, 20, 18, 20, 18, 18, 14, 12, 22),
        META_SHEET: (24, 96),
    }
    for ws in (target, snapshot, meta_sheet):
        ws.auto_filter.ref = ws.dimensions
        for index, width in enumerate(widths[ws.title], start=1):
            ws.column_dimensions[ws.cell(1, index).column_letter].width = width

    # Informational sheets are protected against accidental edits; there is no
    # password and no write path ever trusts CURRENT_SNAPSHOT anyway.
    snapshot.protection.sheet = True
    meta_sheet.protection.sheet = True

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return {
        "schema_version": 1,
        "file": destination.name,
        "catalog_year": year,
        "base_release_id": release_id,
        "baseline_hash": baseline,
        "target_model_ids": list(selected),
        "target_rows": len(target_rows),
        "snapshot_rows": len(snapshot_rows),
        "compile_only": True,
    }


__all__ = [
    "META_SHEET",
    "MODE",
    "RetailLineupWorkbookCompileResult",
    "RetailLineupWorkbookError",
    "RetailLineupWorkbookMeta",
    "SNAPSHOT_SHEET",
    "TARGET_SHEET",
    "compile_retail_lineup_workbook",
    "generate_retail_lineup_workbook",
    "read_retail_lineup_workbook",
]
