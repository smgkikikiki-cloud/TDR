from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from vehreg.current_retail import replace_current_retail_set
from vehreg.normalize import trim_identity
from vehreg.retail_lineup_bootstrap import LineupAction, RetailLineupBootstrapError
from vehreg.retail_lineup_workbook import (
    META_SHEET,
    SNAPSHOT_SHEET,
    TARGET_SHEET,
    RetailLineupWorkbookError,
    compile_retail_lineup_workbook,
    generate_retail_lineup_workbook,
    read_retail_lineup_workbook,
)

YEAR = 2026
MODEL = "acme.echo"
GEN = f"{MODEL}.e1"
A = f"{GEN}.trim.a_bev"
B = f"{GEN}.trim.b_bev"


def _write(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _seed(tmp_path: Path) -> Path:
    data = tmp_path / "data"
    _write(data / str(YEAR) / "models" / "acme.json", {
        "brand": {
            "id": "acme", "name_en": "Acme", "name_th": "",
            "brand_segment": "MASS", "brand_origin": "US", "aliases": [],
        },
        "models": [{
            "id": "echo", "name_en": "Echo", "name_th": "",
            "body_type": "SEDAN", "cab_type": "NOT_APPLICABLE",
            "registration_type": "RY1", "market_scope": "CORE", "aliases": [],
            "retail_status": "UNVERIFIED",
            "generations": [{
                "code": "E1", "segment": "C", "seats": 5,
                "launched": "2025-01-01", "ended": None,
                "variants": [{
                    "id": "bev", "name": "Echo BEV", "powertrain": "BEV",
                    "drivetrain": "FWD", "price_thb": None, "import_type": "CBU",
                    "origin_country": "US", "aliases": [], "battery_kwh": 60.0,
                }],
                "trims": [
                    {"id": "a_bev", "name": "A", "powertrain": "BEV",
                     "aliases": [], "source_refs": {}},
                    {"id": "b_bev", "name": "B", "powertrain": "BEV",
                     "aliases": [], "source_refs": {}},
                ],
            }],
        }],
    })
    replace_current_retail_set(
        data_dir=data, year=YEAR, model_id=MODEL, trim_ids=[A, B],
        reviewer="Owner", reviewed_at="2026-09-30", source_ref="", write=True,
    )
    _write(data / str(YEAR) / "market" / "prices" / "echo.json", {
        "prices": [
            {"trim_id": A, "amount_thb": 900000, "price_type": "LIST_PRICE"},
            {"trim_id": B, "amount_thb": 950000, "price_type": "CAMPAIGN_PRICE",
             "campaign_id": "echo-campaign", "option_id": "cash"},
        ]
    })
    _write(data / str(YEAR) / "product" / "comparable_specs" / "facts" / "echo.json", {
        "schema_version": 1,
        "facts": [
            {"fact_id": "one", "trim_id": A, "field_key": "dimensions.length_mm"},
            {"fact_id": "two", "trim_id": A, "field_key": "dimensions.width_mm"},
        ],
    })
    return data


def _generate(tmp_path: Path, data: Path) -> Path:
    path = tmp_path / "lineup.xlsx"
    generate_retail_lineup_workbook(
        path,
        model_ids=[MODEL],
        base_release_id="vehicle-test-release",
        data_dir=data,
        year=YEAR,
        generated_at="2026-09-30T12:00:00+07:00",
        as_of="2026-09-30",
    )
    return path


def _actions(result):
    return {
        item.canonical_trim_id: item.action
        for model in result.plan.models
        for item in model.items
    }


def test_generator_writes_three_sheet_contract_and_current_seed(tmp_path):
    data = _seed(tmp_path)
    path = _generate(tmp_path, data)
    workbook = load_workbook(path, data_only=False)

    assert workbook.sheetnames == [TARGET_SHEET, SNAPSHOT_SHEET, META_SHEET]
    target = workbook[TARGET_SHEET]
    assert [cell.value for cell in target[1]] == [
        "model_id", "generation_id", "trim_name", "powertrain",
        "canonical_trim_id", "notes",
    ]
    assert {target["E2"].value, target["E3"].value} == {A, B}

    snapshot = workbook[SNAPSHOT_SHEET]
    headers = {cell.value: index + 1 for index, cell in enumerate(snapshot[1])}
    rows = {
        snapshot.cell(row, headers["canonical_trim_id"]).value: row
        for row in range(2, snapshot.max_row + 1)
    }
    assert snapshot.cell(rows[A], headers["resolved_current"]).value == "YES"
    assert snapshot.cell(rows[A], headers["price_records"]).value == 1
    assert snapshot.cell(rows[A], headers["spec_facts"]).value == 2
    assert snapshot.cell(rows[B], headers["campaign_price_records"]).value == 1

    meta, rows = read_retail_lineup_workbook(path)
    assert meta.mode == "REPLACE_AND_ARCHIVE"
    assert meta.base_release_id == "vehicle-test-release"
    assert meta.target_model_ids == (MODEL,)
    assert len(rows) == 2


def test_compile_is_preview_only_and_untouched_workbook_is_all_keep(tmp_path):
    data = _seed(tmp_path)
    path = _generate(tmp_path, data)
    current_path = data / str(YEAR) / "market" / "trims" / "current_retail.json"
    before = current_path.read_bytes()

    result = compile_retail_lineup_workbook(path, data_dir=data, expected_year=YEAR)

    assert result.rows_read == 2
    assert result.plan.counts == {"keep": 2, "create": 0, "reactivate": 0, "archive": 0}
    assert current_path.read_bytes() == before
    assert result.as_dict()["compile_only"] is True


def test_edit_target_compiles_keep_create_archive_and_snapshot_is_ignored(tmp_path):
    data = _seed(tmp_path)
    path = _generate(tmp_path, data)
    workbook = load_workbook(path)

    target = workbook[TARGET_SHEET]
    target.delete_rows(2, target.max_row - 1)
    target.append([MODEL, GEN, "B", "BEV", B, ""])
    target.append([MODEL, GEN, "C New", "BEV", "", "owner target"])

    snapshot = workbook[SNAPSHOT_SHEET]
    snapshot.protection.sheet = False
    snapshot["F2"] = "TOTALLY WRONG SNAPSHOT NAME"
    workbook.save(path)

    result = compile_retail_lineup_workbook(path, data_dir=data)
    new_id = trim_identity(GEN, None, "C New", "BEV")
    actions = _actions(result)
    assert actions[B] is LineupAction.KEEP
    assert actions[new_id] is LineupAction.CREATE
    assert actions[A] is LineupAction.ARCHIVE
    assert result.plan.models[0].target_current_trim_ids == tuple(sorted({B, new_id}))


def test_stale_workbook_fails_closed_before_plan(tmp_path):
    data = _seed(tmp_path)
    path = _generate(tmp_path, data)

    replace_current_retail_set(
        data_dir=data, year=YEAR, model_id=MODEL, trim_ids=[A],
        reviewer="Owner", reviewed_at="2026-09-30", source_ref="", write=True,
    )

    with pytest.raises(RetailLineupBootstrapError, match="STALE_BASELINE"):
        compile_retail_lineup_workbook(path, data_dir=data)


def test_empty_model_target_never_means_withdraw(tmp_path):
    data = _seed(tmp_path)
    path = _generate(tmp_path, data)
    workbook = load_workbook(path)
    target = workbook[TARGET_SHEET]
    target.delete_rows(2, target.max_row - 1)
    workbook.save(path)

    with pytest.raises(RetailLineupWorkbookError, match="at least one target trim"):
        compile_retail_lineup_workbook(path, data_dir=data)


def test_target_formula_and_unknown_header_are_rejected(tmp_path):
    data = _seed(tmp_path)
    formula_path = _generate(tmp_path, data)
    workbook = load_workbook(formula_path)
    workbook[TARGET_SHEET]["D2"] = '=UPPER("bev")'
    workbook.save(formula_path)
    with pytest.raises(RetailLineupWorkbookError, match="formulas are not allowed"):
        compile_retail_lineup_workbook(formula_path, data_dir=data)

    header_path = tmp_path / "bad-header.xlsx"
    generate_retail_lineup_workbook(
        header_path, model_ids=[MODEL], base_release_id="vehicle-test-release",
        data_dir=data, year=YEAR, generated_at="2026-09-30T12:00:00+07:00",
        as_of="2026-09-30",
    )
    workbook = load_workbook(header_path)
    workbook[TARGET_SHEET]["G1"] = "surprise"
    workbook.save(header_path)
    with pytest.raises(RetailLineupWorkbookError, match="headers must be exactly"):
        compile_retail_lineup_workbook(header_path, data_dir=data)


def test_meta_target_model_ids_prevents_silent_row_scope_change(tmp_path):
    data = _seed(tmp_path)
    path = _generate(tmp_path, data)
    workbook = load_workbook(path)
    meta = workbook[META_SHEET]
    meta.protection.sheet = False
    for row in range(2, meta.max_row + 1):
        if meta.cell(row, 1).value == "target_model_ids":
            meta.cell(row, 2).value = json.dumps([MODEL, "acme.ghost"])
            break
    workbook.save(path)

    with pytest.raises(RetailLineupWorkbookError, match="missing rows"):
        compile_retail_lineup_workbook(path, data_dir=data)
