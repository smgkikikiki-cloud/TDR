"""Export a fillable workbook for direct canonical vehicle-spec editing.

Examples:

    python -m tools.export_vehicle_spec_template --model-id toyota.corolla-cross --out /tmp/specs.xlsx
    python -m tools.export_vehicle_spec_template --trim-id <canonical-trim-id> --out /tmp/specs.xlsx

The SPECS sheet contains exact canonical identities plus deterministic machine
headers. Every current registry field is represented. Fields whose meaning
requires qualifier context are emitted as explicit fixed columns instead of an
ambiguous bare value column.
"""

from __future__ import annotations

import argparse
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from vehreg.catalog import Catalog, DATA_DIR, DEFAULT_YEAR
from vehreg.comparable_specs import SpecRegistry
from vehreg.spec_excel import CORE_ONLY_HEADERS, HEADER_ALIASES, IDENTITY_COLUMN

DISPLAY_COLUMNS = ("brand", "model", "generation", "trim", "powertrain")

# Fixed machine headers for qualified registry fields that do not already have
# a short alias in vehreg.spec_excel. The qualifier is part of the column name,
# so a value such as 300 can never be silently guessed as NEDC/CLTC/WLTP.
# Multiple columns are intentional where the same canonical field can carry
# multiple simultaneously valid contexts.
QUALIFIED_TEMPLATE_HEADERS: dict[str, tuple[str, ...]] = {
    "vehicle.curb_weight_kg": (
        "vehicle.curb_weight_kg__measurement_basis=DECLARED",
    ),
    "vehicle.cargo_volume_l": (
        "vehicle.cargo_volume_l__measurement_basis=VDA__seat_configuration=SEATS_UP",
        "vehicle.cargo_volume_l__measurement_basis=VDA__seat_configuration=REAR_FOLDED",
    ),
    "vehicle.turning_radius_m": (
        "vehicle.turning_radius_m__measurement_basis=CURB_TO_CURB",
    ),
    "powertrain.max_power_kw": (
        "powertrain.max_power_kw__output_scope=SYSTEM__rating_basis=DECLARED",
        "powertrain.max_power_kw__output_scope=ENGINE__rating_basis=DECLARED",
        "powertrain.max_power_kw__output_scope=FRONT_MOTOR__rating_basis=DECLARED",
        "powertrain.max_power_kw__output_scope=REAR_MOTOR__rating_basis=DECLARED",
    ),
    "powertrain.max_torque_nm": (
        "powertrain.max_torque_nm__output_scope=SYSTEM__rating_basis=DECLARED",
        "powertrain.max_torque_nm__output_scope=ENGINE__rating_basis=DECLARED",
        "powertrain.max_torque_nm__output_scope=FRONT_MOTOR__rating_basis=DECLARED",
        "powertrain.max_torque_nm__output_scope=REAR_MOTOR__rating_basis=DECLARED",
    ),
    "performance.acceleration_0_100_s": (
        "performance.acceleration_0_100_s__measurement_basis=DECLARED",
    ),
    "performance.top_speed_kmh": (
        "performance.top_speed_kmh__measurement_basis=DECLARED",
    ),
    "ev.energy_consumption_wh_km": tuple(
        f"ev.energy_consumption_wh_km__measurement_basis={basis}__range_scope={scope}"
        for scope in ("FULL", "ELECTRIC_ONLY")
        for basis in ("NEDC", "CLTC", "WLTP", "WLTC", "EPA")
    ),
    "charging.dc_max_kw": (
        "charging.dc_max_kw__rating_basis=DECLARED",
        "charging.dc_max_kw__rating_basis=PEAK",
    ),
    "emissions.co2_g_km": tuple(
        f"emissions.co2_g_km__measurement_basis={basis}"
        for basis in ("ECO_STICKER_TH", "NEDC", "WLTP", "WLTC")
    ),
    "safety.ncap_stars": tuple(
        f"safety.ncap_stars__program={program}"
        for program in ("ASEAN_NCAP", "EURO_NCAP", "ANCAP", "C_NCAP")
    ),
    "efficiency.fuel_consumption_l_100km": tuple(
        f"efficiency.fuel_consumption_l_100km__measurement_basis={basis}"
        for basis in ("ECO_STICKER_TH", "NEDC", "WLTP", "WLTC")
    ),
    "efficiency.fuel_consumption_urban_l_100km": tuple(
        f"efficiency.fuel_consumption_urban_l_100km__measurement_basis={basis}"
        for basis in ("ECO_STICKER_TH", "NEDC", "WLTP", "WLTC")
    ),
    "efficiency.fuel_consumption_extra_urban_l_100km": tuple(
        f"efficiency.fuel_consumption_extra_urban_l_100km__measurement_basis={basis}"
        for basis in ("ECO_STICKER_TH", "NEDC", "WLTP", "WLTC")
    ),
}


def _selected_trims(catalog: Catalog, *, model_id: str | None, trim_id: str | None):
    if trim_id:
        trim = catalog.trims.get(trim_id)
        if trim is None:
            raise ValueError(f"unknown canonical trim id {trim_id!r}")
        return [trim]
    if model_id:
        if model_id not in catalog.models:
            raise ValueError(f"unknown canonical model id {model_id!r}")
        return sorted(catalog.trims_of(model_id), key=lambda row: row.id)
    return sorted(catalog.trims.values(), key=lambda row: row.id)


def _ready_columns_by_field() -> dict[str, list[str]]:
    out: dict[str, list[str]] = defaultdict(list)
    for header, (field_key, _qualifiers) in HEADER_ALIASES.items():
        out[field_key].append(header)
    for field_key, headers in QUALIFIED_TEMPLATE_HEADERS.items():
        out[field_key].extend(headers)
    return {key: list(dict.fromkeys(values)) for key, values in out.items()}


def workbook_headers(registry: SpecRegistry) -> list[str]:
    ready = _ready_columns_by_field()
    columns = [IDENTITY_COLUMN, *DISPLAY_COLUMNS, *sorted(CORE_ONLY_HEADERS)]
    for key, definition in sorted(registry.fields.items()):
        if definition.comparison_qualifiers:
            qualified = ready.get(key, [])
            if not qualified:
                raise ValueError(
                    f"qualified registry field {key!r} has no deterministic template column")
            columns.extend(qualified)
        else:
            columns.append(key)
    # One header exactly once, preserving the stable order above.
    return list(dict.fromkeys(columns))


def export_template(
    out: Path,
    *,
    year: int = DEFAULT_YEAR,
    model_id: str | None = None,
    trim_id: str | None = None,
) -> dict:
    catalog = Catalog.load(DATA_DIR, year)
    registry = SpecRegistry.load(DATA_DIR, year)
    trims = _selected_trims(catalog, model_id=model_id, trim_id=trim_id)
    headers = workbook_headers(registry)
    ready = _ready_columns_by_field()

    wb = Workbook()
    ws = wb.active
    ws.title = "SPECS"
    ws.append(headers)
    for trim in trims:
        generation = catalog.generations[trim.generation_id]
        model = catalog.models[generation.model_id]
        brand = catalog.brands[model.brand_id]
        identity = {
            IDENTITY_COLUMN: trim.id,
            "brand": brand.name_en,
            "model": model.name_en,
            "generation": generation.code,
            "trim": trim.name,
            "powertrain": trim.powertrain.value,
        }
        ws.append([identity.get(header, "") for header in headers])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    ws.freeze_panes = "G2"
    ws.auto_filter.ref = ws.dimensions
    for index, header in enumerate(headers, start=1):
        width = 18
        if header == IDENTITY_COLUMN:
            width = 48
        elif header in DISPLAY_COLUMNS:
            width = 22
        else:
            width = min(52, max(14, len(header) + 2))
        ws.column_dimensions[get_column_letter(index)].width = width

    dictionary = wb.create_sheet("FIELD_DICTIONARY")
    dictionary_headers = [
        "field_key", "label_en", "label_th", "value_type", "unit",
        "applicable_powertrains", "qualifiers", "ready_columns", "custom_header_rule",
    ]
    dictionary.append(dictionary_headers)
    for key, definition in sorted(registry.fields.items()):
        ready_columns = ready.get(key, []) if definition.comparison_qualifiers else [key]
        rule = (
            f"{key}__qualifier=value" if definition.comparison_qualifiers else key
        )
        dictionary.append([
            key,
            definition.label_en,
            definition.label_th,
            definition.value_type.value,
            definition.canonical_unit,
            ",".join(definition.applicable_powertrains),
            ",".join(definition.comparison_qualifiers),
            ",".join(ready_columns),
            rule,
        ])
    for key, core_field in sorted(CORE_ONLY_HEADERS.items()):
        dictionary.append([
            key, f"MarketTrim core: {core_field}", "", "CORE", "", "", "", key, key,
        ])
    for cell in dictionary[1]:
        cell.fill = header_fill
        cell.font = header_font
    dictionary.freeze_panes = "A2"
    dictionary.auto_filter.ref = dictionary.dimensions
    for index, width in enumerate((36, 30, 30, 14, 12, 30, 34, 70, 70), start=1):
        dictionary.column_dimensions[get_column_letter(index)].width = width

    notes = wb.create_sheet("README")
    notes.append(["Vehicle Spec Excel Import"])
    notes.append(["1", "Fill only cells you want to change. Blank = no-op."])
    notes.append(["2", "Do not edit canonical_trim_id. It is the only placement identity."])
    notes.append(["3", "Use YES/NO (or TRUE/FALSE, 1/0) for boolean fields."])
    notes.append(["4", "Use UNKNOWN / NOT_AVAILABLE / NOT_APPLICABLE when needed. '-' is rejected."])
    notes.append(["5", "Qualifier context is already encoded in SPECS headers (for example NEDC/WLTP/CLTC). Do not guess or strip it."])
    notes["A1"].font = Font(bold=True, size=14)
    notes.column_dimensions["A"].width = 10
    notes.column_dimensions["B"].width = 110

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return {
        "output": str(out),
        "trims": len(trims),
        "data_columns": len(headers),
        "registry_fields": len(registry.fields),
    }


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--year", type=int, default=DEFAULT_YEAR)
    group = parser.add_mutually_exclusive_group()
    group.add_argument("--model-id")
    group.add_argument("--trim-id")
    args = parser.parse_args(argv)
    result = export_template(
        args.out,
        year=args.year,
        model_id=args.model_id,
        trim_id=args.trim_id,
    )
    print(result)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
